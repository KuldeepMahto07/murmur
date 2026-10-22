"""Transcript exporters: TXT, SRT, VTT, CSV, JSON, DOCX, PDF."""
import csv
import io
import json
from dataclasses import dataclass
from pathlib import Path

FORMATS = {
    "txt": ("text/plain; charset=utf-8", "txt"),
    "srt": ("application/x-subrip; charset=utf-8", "srt"),
    "vtt": ("text/vtt; charset=utf-8", "vtt"),
    "csv": ("text/csv; charset=utf-8", "csv"),
    "json": ("application/json", "json"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "pdf": ("application/pdf", "pdf"),
}


@dataclass
class ExportOptions:
    timestamps: bool = True
    speakers: bool = True
    max_chars: int = 84  # subtitle cue length (two lines of ~42)
    max_seconds: float = 6.0


def hms(t: float, sep: str = ".", ms: bool = False) -> str:
    t = max(t, 0)
    h, rem = divmod(int(t), 3600)
    m, s = divmod(rem, 60)
    base = f"{h:02d}:{m:02d}:{s:02d}"
    return f"{base}{sep}{int(round((t - int(t)) * 1000)) % 1000:03d}" if ms else base


def _name(seg: dict, speakers: dict) -> str | None:
    spk = seg.get("speaker")
    return speakers.get(spk, _default_name(spk)) if spk else None


def _default_name(spk: str) -> str:
    try:
        return f"Speaker {int(spk.split('_')[-1]) + 1}"
    except ValueError:
        return spk


def paragraphs(transcript: dict, speakers: dict) -> list[dict]:
    """Group consecutive segments by speaker (or by pauses when there are no speakers)."""
    out: list[dict] = []
    for seg in transcript["segments"]:
        name = _name(seg, speakers)
        if out and out[-1]["speaker"] == name and (
                name is not None or seg["start"] - out[-1]["end"] < 2.0) and len(out[-1]["text"]) < 1200:
            out[-1]["text"] += " " + seg["text"]
            out[-1]["end"] = seg["end"]
        else:
            out.append({"start": seg["start"], "end": seg["end"], "speaker": name, "text": seg["text"]})
    return out


def cues(transcript: dict, speakers: dict, opt: ExportOptions) -> list[dict]:
    """Split segments into subtitle-sized cues using word timings when available."""
    result = []
    for seg in transcript["segments"]:
        name = _name(seg, speakers) if opt.speakers else None
        words = seg.get("words") or []
        if not words or (len(seg["text"]) <= opt.max_chars and seg["end"] - seg["start"] <= opt.max_seconds):
            result.append({"start": seg["start"], "end": seg["end"], "text": seg["text"], "speaker": name})
            continue
        cur, text = [], ""
        for w in words:
            if cur and (len(text + w["word"]) > opt.max_chars or w["end"] - cur[0]["start"] > opt.max_seconds):
                result.append({"start": cur[0]["start"], "end": cur[-1]["end"], "text": text.strip(), "speaker": name})
                cur, text = [], ""
            cur.append(w)
            text += w["word"]
        if cur:
            result.append({"start": cur[0]["start"], "end": cur[-1]["end"], "text": text.strip(), "speaker": name})
    return result


def _wrap(text: str, width: int = 42) -> str:
    if len(text) <= width:
        return text
    cut = text.rfind(" ", 0, len(text) // 2 + 10)
    return text if cut <= 0 else text[:cut] + "\n" + text[cut + 1:]


def to_txt(t, speakers, opt):
    lines = []
    for p in paragraphs(t, speakers):
        head = []
        if opt.timestamps:
            head.append(f"[{hms(p['start'])}]")
        if opt.speakers and p["speaker"]:
            head.append(f"{p['speaker']}:")
        lines.append((" ".join(head) + "\n" if head else "") + p["text"])
    return "\n\n".join(lines) + "\n"


def to_srt(t, speakers, opt):
    blocks = []
    for i, c in enumerate(cues(t, speakers, opt), 1):
        text = f"[{c['speaker']}] {c['text']}" if c["speaker"] else c["text"]
        blocks.append(f"{i}\n{hms(c['start'], ',', True)} --> {hms(c['end'], ',', True)}\n{_wrap(text)}")
    return "\n\n".join(blocks) + "\n"


def to_vtt(t, speakers, opt):
    blocks = ["WEBVTT"]
    for c in cues(t, speakers, opt):
        text = f"<v {c['speaker']}>{c['text']}" if c["speaker"] else c["text"]
        blocks.append(f"{hms(c['start'], '.', True)} --> {hms(c['end'], '.', True)}\n{_wrap(text)}")
    return "\n\n".join(blocks) + "\n"


def to_csv(t, speakers, opt):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["start", "end", "speaker", "text"])
    for seg in t["segments"]:
        w.writerow([hms(seg["start"], ".", True), hms(seg["end"], ".", True), _name(seg, speakers) or "", seg["text"]])
    return "\ufeff" + buf.getvalue()  # BOM so Excel reads UTF-8


def to_json(t, speakers, opt):
    data = dict(t)
    data["speaker_names"] = {k: speakers.get(k, _default_name(k)) for k in
                             sorted({s["speaker"] for s in t["segments"] if s.get("speaker")})}
    return json.dumps(data, ensure_ascii=False, indent=2)


def to_docx(t, speakers, opt, title: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_heading(title, level=1)
    for p in paragraphs(t, speakers):
        para = doc.add_paragraph()
        if opt.speakers and p["speaker"]:
            r = para.add_run(p["speaker"])
            r.bold = True
        if opt.timestamps:
            r = para.add_run(("  " if opt.speakers and p["speaker"] else "") + hms(p["start"]))
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        if para.runs:
            para.add_run().add_break()
        para.add_run(p["text"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for name, path in (("Body", "C:/Windows/Fonts/arial.ttf"), ("Body", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
                       ("Body", "/Library/Fonts/Arial Unicode.ttf")):
        if Path(path).exists():
            if "Body" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, path))
            return name
    return "Helvetica"


def to_pdf(t, speakers, opt, title: str) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    font = _pdf_font()
    h = ParagraphStyle("h", fontName=font, fontSize=18, leading=22, spaceAfter=10)
    meta = ParagraphStyle("m", fontName=font, fontSize=9, leading=12, textColor="#777777")
    body = ParagraphStyle("b", fontName=font, fontSize=11, leading=15, spaceAfter=8)
    story = [Paragraph(escape(title), h)]
    for p in paragraphs(t, speakers):
        head = []
        if opt.speakers and p["speaker"]:
            head.append(f"<b>{escape(p['speaker'])}</b>")
        if opt.timestamps:
            head.append(hms(p["start"]))
        if head:
            story.append(Paragraph("  ·  ".join(head), meta))
        story.append(Paragraph(escape(p["text"]), body))
    story.append(Spacer(1, 4 * mm))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title=title, leftMargin=20 * mm, rightMargin=20 * mm,
                      topMargin=18 * mm, bottomMargin=18 * mm).build(story)
    return buf.getvalue()


def render(fmt: str, transcript: dict, speakers: dict, opt: ExportOptions, title: str) -> bytes:
    if fmt == "docx":
        return to_docx(transcript, speakers, opt, title)
    if fmt == "pdf":
        return to_pdf(transcript, speakers, opt, title)
    fn = {"txt": to_txt, "srt": to_srt, "vtt": to_vtt, "csv": to_csv, "json": to_json}[fmt]
    return fn(transcript, speakers, opt).encode("utf-8")
